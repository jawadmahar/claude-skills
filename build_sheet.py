import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Daily Takings"

# ---- config ----
labels = ["Date", "Gross £", "Card Machine", "Worldpay", "Total card",
          "Expenses", "Red box", "Cash 2 bank", "Till Total"]
# day columns per block (11 wide grid; aligns vertically). 11+11+9 = 31 days
blocks = [list(range(1, 12)), list(range(12, 23)), list(range(23, 32))]
NDAYS_COLS = 11            # B..L
last_col = 1 + NDAYS_COLS  # col index of L

# ---- styles ----
thin = Side(style="thin", color="000000")
med  = Side(style="medium", color="000000")
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)
title_font = Font(name="Calibri", size=18, bold=True)
sub_font   = Font(name="Calibri", size=11, bold=True)
label_font = Font(name="Calibri", size=10, bold=True)
day_font   = Font(name="Calibri", size=11, bold=True)
center = Alignment(horizontal="center", vertical="center")
left   = Alignment(horizontal="left", vertical="center", indent=1)
hdr_fill   = PatternFill("solid", fgColor="D9E1F2")   # light blue for day header
label_fill = PatternFill("solid", fgColor="F2F2F2")   # light grey for labels
date_fill  = PatternFill("solid", fgColor="BDD7EE")   # stronger blue for Date row

# ---- title block ----
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
c = ws.cell(row=1, column=1, value="ENGINO STEM CENTRE")
c.font = title_font; c.alignment = center

ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=4)
c = ws.cell(row=2, column=1, value="Daily Takings  —  Month / Year:")
c.font = sub_font; c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
# blank fill-in line for the month
ws.merge_cells(start_row=2, start_column=5, end_row=2, end_column=last_col)
mc = ws.cell(row=2, column=5)
mc.border = Border(bottom=med)

r = 4  # first data row
for blk in blocks:
    # header row: "Date" label + day numbers
    lc = ws.cell(row=r, column=1, value="Date")
    lc.font = label_font; lc.alignment = left; lc.fill = date_fill; lc.border = border_all
    for i in range(NDAYS_COLS):
        col = 2 + i
        cell = ws.cell(row=r, column=col)
        if i < len(blk):
            cell.value = blk[i]
            cell.font = day_font; cell.alignment = center
            cell.fill = hdr_fill; cell.border = border_all
        # cells beyond the block's days stay empty/unbordered
    # metric rows
    for label in labels[1:]:
        rr = r + labels.index(label)
        lc = ws.cell(row=rr, column=1, value=label)
        lc.font = label_font; lc.alignment = left; lc.fill = label_fill; lc.border = border_all
        for i in range(NDAYS_COLS):
            col = 2 + i
            cell = ws.cell(row=rr, column=col)
            if i < len(blk):
                cell.border = border_all
                cell.alignment = center
    r += len(labels) + 1   # blank spacer row between blocks

# ---- column widths ----
ws.column_dimensions["A"].width = 15
for i in range(NDAYS_COLS):
    ws.column_dimensions[get_column_letter(2 + i)].width = 8.5

# ---- row heights ----
ws.row_dimensions[1].height = 26
ws.row_dimensions[2].height = 20
for rr in range(4, r):
    if ws.row_dimensions[rr].height is None:
        ws.row_dimensions[rr].height = 22

# ---- print setup: A4 landscape, fit to one page wide ----
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins.left = 0.3
ws.page_margins.right = 0.3
ws.page_margins.top = 0.4
ws.page_margins.bottom = 0.4
ws.print_area = f"A1:{get_column_letter(last_col)}{r-1}"
ws.sheet_view.showGridLines = False

out = "/home/user/claude-skills/output/ENGINO_STEM_CENTRE_daily_takings_blank.xlsx"
wb.save(out)
print("saved", out)
